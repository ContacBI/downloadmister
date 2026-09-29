"""Compara o Aging oficial (xlsm) com a listagem de notas de dezembro.

Uso: python3 comparar_oficial.py oficial.xlsm clientes-nota.csv saida.xlsx

Vínculo por (cliente resumido, nota). O oficial vem por parcela, então é
somado por nota antes de comparar parcelas, emissão, valor, vencido e a vencer.
Abas: Resumo, Matches exatos, Só no oficial, Só em dezembro,
Valores diferentes e Por cliente.
"""
import csv
import sys
from collections import defaultdict

import openpyxl

TOL = 0.01


def num(s):
    s = s.strip().replace("\xa0", "")
    if s in ("", "-"):
        return 0.0
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    return float(s)


def nf(s):
    return str(s).strip().lstrip("0") or "0"


def ler_oficial(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = list(wb["Aging analítico - Contabil"].iter_rows(values_only=True))
    of = {}
    for r in rows[6:]:
        if not r[1] or r[1] == "Total" or not r[2]:
            continue
        k = (r[1].strip(), nf(r[2]))
        a = of.setdefault(k, {"parc": 0, "emissao": r[3].strftime("%d/%m/%Y"), "valor": 0.0, "venc": 0.0, "avenc": 0.0})
        a["parc"] += 1
        a["valor"] += r[5] or 0
        a["venc"] += r[14] or 0
        a["avenc"] += r[20] or 0
    return of


def ler_dezembro(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))[1:]
    return {
        (r[0].strip(), nf(r[3])): {"parc": int(r[5]), "emissao": r[4], "valor": num(r[6]),
                                   "venc": num(r[7]), "avenc": num(r[8])}
        for r in rows
    }


def igual(d, o):
    return (d["parc"] == o["parc"] and d["emissao"] == o["emissao"]
            and all(abs(d[c] - o[c]) <= TOL for c in ("valor", "venc", "avenc")))


def main(oficial, nota, saida):
    of, dz = ler_oficial(oficial), ler_dezembro(nota)
    so_of = sorted(k for k in of if k not in dz)
    so_dz = sorted(k for k in dz if k not in of)
    ambos = sorted(k for k in of if k in dz)
    exatos = [k for k in ambos if igual(dz[k], of[k])]
    dif = [k for k in ambos if not igual(dz[k], of[k])]
    r2 = lambda x: round(x, 2)
    tot = lambda d, c: r2(sum(a[c] for a in d.values()))

    por_cli = defaultdict(lambda: [0.0, 0.0])
    for (c, _), a in of.items():
        por_cli[c][0] += a["valor"]
    for (c, _), a in dz.items():
        por_cli[c][1] += a["valor"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["", "Dezembro (nossa lista)", "Oficial", "Diferença"])
    ws.append(["Notas", len(dz), len(of), len(of) - len(dz)])
    for nome, c in (("Valor", "valor"), ("Vencido", "venc"), ("A vencer", "avenc")):
        ws.append([nome, tot(dz, c), tot(of, c), r2(tot(of, c) - tot(dz, c))])
    ws.append([])
    ws.append(["Matches exatos", len(exatos)])
    ws.append(["Só no oficial", len(so_of)])
    ws.append(["Só em dezembro", len(so_dz)])
    ws.append(["Nos dois com valores diferentes", len(dif)])

    ws = wb.create_sheet("Matches exatos")
    ws.append(["Cliente", "Nota", "Emissão", "Parcelas", "Valor", "Vencido", "A vencer"])
    for k in exatos:
        a = of[k]
        ws.append([k[0], k[1], a["emissao"], a["parc"], r2(a["valor"]), r2(a["venc"]), r2(a["avenc"])])

    for titulo, chaves, base in (("Só no oficial", so_of, of), ("Só em dezembro", so_dz, dz)):
        ws = wb.create_sheet(titulo)
        ws.append(["Cliente", "Nota", "Emissão", "Parcelas", "Valor", "Vencido", "A vencer"])
        for k in chaves:
            a = base[k]
            ws.append([k[0], k[1], a["emissao"], a["parc"], r2(a["valor"]), r2(a["venc"]), r2(a["avenc"])])

    ws = wb.create_sheet("Valores diferentes")
    ws.append(["Cliente", "Nota", "Emissão dez", "Emissão oficial", "Parc. dez", "Parc. oficial",
               "Valor dez", "Valor oficial", "Vencido dez", "Vencido oficial", "A vencer dez", "A vencer oficial"])
    for k in dif:
        d, o = dz[k], of[k]
        ws.append([k[0], k[1], d["emissao"], o["emissao"], d["parc"], o["parc"], r2(d["valor"]),
                   r2(o["valor"]), r2(d["venc"]), r2(o["venc"]), r2(d["avenc"]), r2(o["avenc"])])

    ws = wb.create_sheet("Por cliente")
    ws.append(["Cliente", "Valor dez", "Valor oficial", "Diferença"])
    for c, (o, d) in sorted(por_cli.items(), key=lambda kv: -abs(kv[1][0] - kv[1][1])):
        if abs(o - d) > TOL:
            ws.append([c, r2(d), r2(o), r2(o - d)])

    for w in wb.worksheets:
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = max(len(str(c.value or "")) for c in col[:60]) + 2
    wb.save(saida)

    print(f"dez: {len(dz)} notas, valor {tot(dz,'valor'):,.2f} | oficial: {len(of)} notas, valor {tot(of,'valor'):,.2f}")
    print(f"exatos {len(exatos)} | só oficial {len(so_of)} | só dezembro {len(so_dz)} | diferentes {len(dif)}")


if __name__ == "__main__":
    main(*sys.argv[1:4])

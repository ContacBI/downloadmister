"""Compara o balancete (Domínio) com o relatório de clientes, por Cód Domínio.

Uso: python3 comparar_conciliacao.py Balancete.xlsx clientes-AAAA-MM-cli.csv saida.xlsx

- Balancete: salve o .xls do Domínio como .xlsx (Excel: Salvar como). Vale a conta 1.1.20.100.* (clientes);
  a coluna "Código" do balancete é o Cód Domínio.
- Relatório: o CSV de clientes (vários clientes podem dividir o mesmo Cód Domínio; eles são somados).
- Entradas: débito do balancete x "Vendas (período)" do relatório.
- Saídas:   crédito do balancete x ("Recebimento (período)" + "Diferença") do relatório.
- Saldo final: "Saldo Atual" do balancete x "Saldo final" do relatório.
"""
import collections
import csv
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TOL = 0.01
r2 = lambda v: round(v + 0.0, 2)


def num(s):
    s = str(s).strip()
    if s in ("", "-"):
        return 0.0
    return float(s.replace(".", "").replace(",", ".")) if "," in s and "." in s else float(s.replace(",", "."))


def carregar(bal, rel):
    ws = openpyxl.load_workbook(bal, data_only=True).active
    periodo = ""
    B = {}
    for r in ws.iter_rows(values_only=True):
        if r[0] == "Período:":
            periodo = str(r[1])
        if isinstance(r[0], (int, float)) and str(r[1] or "").startswith("1.1.20.100"):
            B[str(int(r[0]))] = {"nome": str(r[3]).strip(), "ant": r[7] or 0, "deb": r[9] or 0, "cred": r[11] or 0, "atual": r[13] or 0}
    with open(rel, encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.reader(f, delimiter=";"))[1:]
    R = collections.defaultdict(lambda: {"nomes": [], "ant": 0.0, "ven": 0.0, "sai": 0.0, "fin": 0.0})
    for r in linhas:
        x = R[r[2]]
        x["nomes"].append(r[1])
        x["ant"] += num(r[5]); x["ven"] += num(r[6])
        x["sai"] += -(num(r[7]) + num(r[8]))  # recebimento + diferença (vêm negativos no relatório)
        x["fin"] += num(r[9])
    return B, R, periodo


def montar(B, R):
    linhas = []
    for k in sorted(set(B) | set(R), key=int):
        b, r = B.get(k), R.get(k)
        l = {"cod": k, "nome": b["nome"] if b else "", "clientes": "; ".join(dict.fromkeys(r["nomes"])) if r else "",
             "ant_r": r2(r["ant"]) if r else None, "ant_b": r2(b["ant"]) if b else None,
             "ent_r": r2(r["ven"]) if r else None, "ent_b": r2(b["deb"]) if b else None,
             "sai_r": r2(r["sai"]) if r else None, "sai_b": r2(b["cred"]) if b else None,
             "fin_r": r2(r["fin"]) if r else None, "fin_b": r2(b["atual"]) if b else None}
        z = lambda x: x or 0
        l["d_ant"] = r2(z(l["ant_b"]) - z(l["ant_r"])); l["d_ent"] = r2(z(l["ent_b"]) - z(l["ent_r"]))
        l["d_sai"] = r2(z(l["sai_b"]) - z(l["sai_r"])); l["d_fin"] = r2(z(l["fin_b"]) - z(l["fin_r"]))
        l["origem"] = "nos dois" if b and r else ("só no balancete" if b else "só no relatório")
        l["dif"] = any(abs(l[c]) > TOL for c in ("d_ant", "d_ent", "d_sai", "d_fin"))
        l["obs"] = ""
        linhas.append(l)
    # observações automáticas
    for l in linhas:
        if not l["dif"]:
            l["situacao"] = "Confere"
            continue
        partes = []
        if abs(l["d_fin"]) > TOL: partes.append("saldo final")
        if abs(l["d_ent"]) > TOL: partes.append("entradas")
        if abs(l["d_sai"]) > TOL: partes.append("saídas")
        if abs(l["d_ant"]) > TOL: partes.append("saldo anterior")
        l["situacao"] = "Diverge: " + ", ".join(partes)
        if abs(l["d_fin"]) <= TOL and abs(l["d_ent"]) > TOL and abs(l["d_ent"] - l["d_sai"]) <= TOL:
            l["obs"] = "Entrada e saída iguais só no balancete/relatório: não muda o saldo final."
    com_fin = [l for l in linhas if abs(l["d_fin"]) > TOL]
    for l in com_fin:  # pares que se anulam: lançamento provavelmente em conta trocada
        for o in com_fin:
            if o is not l and abs(l["d_fin"] + o["d_fin"]) <= TOL:
                l["obs"] = f"Possível lançamento em conta trocada: o oposto ({r2(o['d_fin']):,.2f}) está em {o['cod']} {o['nome'][:30]}."
                break
    return linhas


def escrever(linhas, periodo, saida):
    azul = PatternFill("solid", fgColor="1F3864"); zebra = PatternFill("solid", fgColor="F2F6FC")
    vermelho = PatternFill("solid", fgColor="FCE4E4"); verde = PatternFill("solid", fgColor="E2F0D9")
    fino = Side(style="thin", color="D9D9D9"); borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    money = '#,##0.00;[Red]-#,##0.00;"-"'
    wb = openpyxl.Workbook()

    tot = lambda campo: r2(sum(l[campo] or 0 for l in linhas))
    ws = wb.active; ws.title = "Resumo"; ws.sheet_view.showGridLines = False
    ws["A1"] = "Conciliação de clientes: balancete x relatório"; ws["A1"].font = Font(bold=True, size=16, color="1F3864")
    ws["A2"] = f"Período do balancete: {periodo}"; ws["A2"].font = Font(size=10, color="595959")
    ws.append([]); ws.append(["", "Relatório", "Balancete", "Diferença (balancete - relatório)"])
    for c in ws[4]: c.fill, c.font, c.border = azul, Font(bold=True, color="FFFFFF"), borda; c.alignment = Alignment(horizontal="center", wrap_text=True)
    for nome, a, b in (("Saldo anterior", "ant_r", "ant_b"), ("Entradas (vendas x débito)", "ent_r", "ent_b"), ("Saídas (receb.+dif. x crédito)", "sai_r", "sai_b"), ("Saldo final", "fin_r", "fin_b")):
        ws.append([nome, tot(a), tot(b), r2(tot(b) - tot(a))])
    for row in ws.iter_rows(min_row=5, max_row=8):
        for i, c in enumerate(row):
            c.border = borda
            if i: c.number_format = money
    ws.append([])
    ws.append(["Clientes (Cód Domínio) no balancete", sum(1 for l in linhas if l["origem"] != "só no relatório")])
    ws.append(["Clientes no relatório", sum(1 for l in linhas if l["origem"] != "só no balancete")])
    ws.append(["Conferem em tudo", sum(1 for l in linhas if not l["dif"])])
    ws.append(["Divergem", sum(1 for l in linhas if l["dif"] and l["origem"] == "nos dois")])
    ws.append(["Só no balancete (com diferença)", sum(1 for l in linhas if l["dif"] and l["origem"] == "só no balancete")])
    ws.append(["Só no relatório (com diferença)", sum(1 for l in linhas if l["dif"] and l["origem"] == "só no relatório")])
    ws.append(["Com diferença no saldo final", sum(1 for l in linhas if abs(l["d_fin"]) > TOL)])
    ws.column_dimensions["A"].width = 38
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    for c in "BCD": ws.column_dimensions[c].width = 22

    cols = [("Cód Domínio", "cod", 12, None), ("Cliente (balancete)", "nome", 40, None), ("Clientes do relatório", "clientes", 40, None),
            ("Saldo ant. relatório", "ant_r", 16, money), ("Saldo ant. balancete", "ant_b", 16, money), ("Dif. saldo ant.", "d_ant", 14, money),
            ("Entradas relatório", "ent_r", 16, money), ("Entradas balancete", "ent_b", 16, money), ("Dif. entradas", "d_ent", 14, money),
            ("Saídas relatório", "sai_r", 16, money), ("Saídas balancete", "sai_b", 16, money), ("Dif. saídas", "d_sai", 14, money),
            ("Saldo final relatório", "fin_r", 17, money), ("Saldo final balancete", "fin_b", 17, money), ("Dif. saldo final", "d_fin", 15, money),
            ("Situação", "situacao", 30, None), ("Observação", "obs", 70, None)]

    def aba(titulo, subset, subtitulo):
        w = wb.create_sheet(titulo); w.sheet_view.showGridLines = False
        w["A1"] = titulo; w["A1"].font = Font(bold=True, size=14, color="1F3864")
        w["A2"] = subtitulo; w["A2"].font = Font(size=10, color="595959")
        w.append([]); w.append([c[0] for c in cols])
        for c in w[4]: c.fill, c.font, c.border = azul, Font(bold=True, color="FFFFFF"), borda; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        w.row_dimensions[4].height = 32
        for i, l in enumerate(subset):
            w.append([l[c[1]] for c in cols])
            for j, c in enumerate(cols):
                cell = w.cell(row=5 + i, column=j + 1); cell.border = borda
                if c[3]: cell.number_format = c[3]
                if i % 2: cell.fill = zebra
                if c[1] in ("d_ant", "d_ent", "d_sai", "d_fin") and l[c[1]] is not None and abs(l[c[1]]) > TOL:
                    cell.fill = vermelho; cell.font = Font(bold=True, color="C00000")
                if c[1] == "situacao": cell.fill = verde if l["situacao"] == "Confere" else vermelho
        for j, c in enumerate(cols): w.column_dimensions[get_column_letter(j + 1)].width = c[2]
        w.freeze_panes = "D5"; w.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + max(len(subset), 1)}"
        w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
        w.print_title_rows = "4:4"

    fin = [l for l in linhas if abs(l["d_fin"]) > TOL]
    aba("Saldo final diferente", sorted(fin, key=lambda l: -abs(l["d_fin"])), f"{len(fin)} clientes com saldo final diferente (soma das diferenças: {r2(sum(l['d_fin'] for l in fin)):,.2f})")
    dif = [l for l in linhas if l["dif"]]
    aba("Todas as divergências", dif, f"{len(dif)} clientes com qualquer diferença (saldo anterior, entradas, saídas ou saldo final)")
    aba("Só no balancete", [l for l in linhas if l["origem"] == "só no balancete" and (abs(l["ent_b"]) > TOL or abs(l["sai_b"]) > TOL or abs(l["fin_b"]) > TOL or abs(l["ant_b"]) > TOL)], "Contas do balancete que não aparecem no relatório de clientes (só com movimento ou saldo)")
    aba("Só no relatório", [l for l in linhas if l["origem"] == "só no relatório"], "Cód Domínio do relatório que não existem no balancete")
    aba("Todos os clientes", linhas, f"{len(linhas)} códigos comparados")
    wb.save(saida)


if __name__ == "__main__":
    B, R, periodo = carregar(sys.argv[1], sys.argv[2])
    L = montar(B, R)
    escrever(L, periodo, sys.argv[3])
    print(f"{len(L)} códigos | divergem: {sum(1 for l in L if l['dif'])} | saldo final diferente: {sum(1 for l in L if abs(l['d_fin']) > TOL)}")
